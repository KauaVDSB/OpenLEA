import datetime

import openpyxl
import pytest
from openpyxl.styles import PatternFill
from openpyxl.utils.datetime import to_excel

from config import COR_FALHA, COR_SUCESSO
from excel_logger import carregar_fila_pacientes, registrar_status, verificar_colunas

HEADERS_COMPLETOS = ["REG", "CONSULTA", "MÉDICO", "DATA DE VALIDADE"]


def _criar_planilha(tmp_path, headers, rows, nome="pacientes.xlsx"):
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    caminho = tmp_path / nome
    wb.save(caminho)
    return str(caminho)


def _pintar_celula(caminho, row_index, col_index, cor):
    wb = openpyxl.load_workbook(caminho)
    sheet = wb.active
    preenchimento = PatternFill(start_color=cor, end_color=cor, fill_type="solid")
    sheet.cell(row=row_index, column=col_index).fill = preenchimento
    wb.save(caminho)
    wb.close()


# ---------------------------------------------------------------------------
# verificar_colunas
# ---------------------------------------------------------------------------

def test_verificar_colunas_todas_presentes(tmp_path):
    caminho = _criar_planilha(tmp_path, HEADERS_COMPLETOS, [["123", "C1", "DR. FULANO", "14082026"]])
    assert verificar_colunas(caminho) == []


def test_verificar_colunas_medico_ausente(tmp_path):
    caminho = _criar_planilha(tmp_path, ["REG", "CONSULTA", "DATA DE VALIDADE"], [["123", "C1", "14082026"]])
    assert verificar_colunas(caminho) == ["medico"]


def test_verificar_colunas_data_ausente(tmp_path):
    caminho = _criar_planilha(tmp_path, ["REG", "CONSULTA", "MÉDICO"], [["123", "C1", "DR. FULANO"]])
    assert verificar_colunas(caminho) == ["data"]


def test_verificar_colunas_ambas_ausentes(tmp_path):
    caminho = _criar_planilha(tmp_path, ["REG", "CONSULTA"], [["123", "C1"]])
    assert verificar_colunas(caminho) == ["medico", "data"]


@pytest.mark.parametrize("alias_medico", ["MÉDICO", "MEDICO"])
def test_verificar_colunas_aceita_alias_medico_sem_acento(tmp_path, alias_medico):
    caminho = _criar_planilha(tmp_path, ["REG", alias_medico, "DATA"], [["123", "DR. FULANO", "14082026"]])
    assert "medico" not in verificar_colunas(caminho)


@pytest.mark.parametrize("alias_data", ["DATA DE VALIDADE", "VALIDADE", "DATA"])
def test_verificar_colunas_aceita_aliases_data(tmp_path, alias_data):
    caminho = _criar_planilha(tmp_path, ["REG", "MÉDICO", alias_data], [["123", "DR. FULANO", "14082026"]])
    assert "data" not in verificar_colunas(caminho)


# ---------------------------------------------------------------------------
# carregar_fila_pacientes
# ---------------------------------------------------------------------------

def test_carregar_fila_le_dados_basicos(tmp_path):
    caminho = _criar_planilha(
        tmp_path, HEADERS_COMPLETOS, [["123", "C1", "DR. FULANO", "14082026"]]
    )
    fila = carregar_fila_pacientes(caminho)
    assert fila == [
        {
            "row_index": 2,
            "reg": "123",
            "consulta": "C1",
            "medico": "DR. FULANO",
            "data_validade": "14082026",
        }
    ]


def test_carregar_fila_pula_linha_sem_reg(tmp_path):
    caminho = _criar_planilha(
        tmp_path,
        HEADERS_COMPLETOS,
        [
            ["123", "C1", "DR. FULANO", "14082026"],
            [None, "C2", "DR. FULANO", "14082026"],
            ["125", "C3", "DR. FULANO", "14082026"],
        ],
    )
    fila = carregar_fila_pacientes(caminho)
    assert [p["reg"] for p in fila] == ["123", "125"]


def test_carregar_fila_pula_linha_ja_colorida(tmp_path):
    caminho = _criar_planilha(
        tmp_path,
        HEADERS_COMPLETOS,
        [
            ["123", "C1", "DR. FULANO", "14082026"],
            ["124", "C2", "DR. FULANO", "14082026"],
        ],
    )
    # Linha 2 (paciente 123) já processada com sucesso
    _pintar_celula(caminho, row_index=2, col_index=1, cor=COR_SUCESSO)

    fila = carregar_fila_pacientes(caminho)
    assert [p["reg"] for p in fila] == ["124"]


def test_carregar_fila_arrasta_medico_e_data_para_linhas_vazias(tmp_path):
    caminho = _criar_planilha(
        tmp_path,
        HEADERS_COMPLETOS,
        [
            ["123", "C1", "DR. FULANO", "14082026"],
            ["124", "C2", None, None],
            ["125", "C3", None, None],
            ["126", "C4", "DR. CICLANO", "15082026"],
            ["127", "C5", None, None],
        ],
    )
    fila = carregar_fila_pacientes(caminho)
    medicos = [p["medico"] for p in fila]
    datas = [p["data_validade"] for p in fila]
    assert medicos == ["DR. FULANO", "DR. FULANO", "DR. FULANO", "DR. CICLANO", "DR. CICLANO"]
    assert datas == ["14082026", "14082026", "14082026", "15082026", "15082026"]


def test_carregar_fila_usa_globais_quando_colunas_ausentes(tmp_path):
    caminho = _criar_planilha(
        tmp_path,
        ["REG", "CONSULTA"],
        [["123", "C1"], ["124", "C2"]],
    )
    fila = carregar_fila_pacientes(caminho, medico_global="DR. GLOBAL", data_global="01012026")
    assert all(p["medico"] == "DR. GLOBAL" for p in fila)
    assert all(p["data_validade"] == "01012026" for p in fila)


def test_carregar_fila_consulta_ausente_retorna_vazio(tmp_path):
    caminho = _criar_planilha(
        tmp_path,
        ["REG", "MÉDICO", "DATA"],
        [["123", "DR. FULANO", "14082026"]],
    )
    fila = carregar_fila_pacientes(caminho)
    assert fila[0]["consulta"] == ""


def test_carregar_fila_converte_data_objeto_datetime(tmp_path):
    caminho = _criar_planilha(
        tmp_path,
        HEADERS_COMPLETOS,
        [["123", "C1", "DR. FULANO", datetime.datetime(2026, 8, 14)]],
    )
    fila = carregar_fila_pacientes(caminho)
    assert fila[0]["data_validade"] == "14082026"


def test_carregar_fila_converte_data_numero_serial_excel(tmp_path):
    serial = to_excel(datetime.datetime(2026, 8, 14))
    caminho = _criar_planilha(
        tmp_path,
        HEADERS_COMPLETOS,
        [["123", "C1", "DR. FULANO", serial]],
    )
    fila = carregar_fila_pacientes(caminho)
    assert fila[0]["data_validade"] == "14082026"


def test_carregar_fila_converte_data_string_com_separadores(tmp_path):
    caminho = _criar_planilha(
        tmp_path,
        HEADERS_COMPLETOS,
        [["123", "C1", "DR. FULANO", "14/08/2026"]],
    )
    fila = carregar_fila_pacientes(caminho)
    assert fila[0]["data_validade"] == "14082026"


def test_carregar_fila_levanta_erro_sem_medico_ou_data(tmp_path):
    caminho = _criar_planilha(
        tmp_path,
        HEADERS_COMPLETOS,
        [["123", "C1", None, None]],
    )
    with pytest.raises(ValueError):
        carregar_fila_pacientes(caminho)


def test_carregar_fila_planilha_vazia_retorna_lista_vazia(tmp_path):
    caminho = _criar_planilha(tmp_path, HEADERS_COMPLETOS, [])
    assert carregar_fila_pacientes(caminho) == []


# ---------------------------------------------------------------------------
# registrar_status
# ---------------------------------------------------------------------------

def test_registrar_status_sucesso_pinta_de_azul(tmp_path):
    caminho = _criar_planilha(tmp_path, HEADERS_COMPLETOS, [["123", "C1", "DR. FULANO", "14082026"]])
    registrar_status(caminho, row_index=2, status="sucesso")

    wb = openpyxl.load_workbook(caminho)
    sheet = wb.active
    assert sheet.cell(row=2, column=1).fill.start_color.rgb == COR_SUCESSO
    wb.close()


def test_registrar_status_erro_pinta_de_vermelho(tmp_path):
    caminho = _criar_planilha(tmp_path, HEADERS_COMPLETOS, [["123", "C1", "DR. FULANO", "14082026"]])
    registrar_status(caminho, row_index=2, status="erro")

    wb = openpyxl.load_workbook(caminho)
    sheet = wb.active
    assert sheet.cell(row=2, column=1).fill.start_color.rgb == COR_FALHA
    wb.close()


def test_registrar_status_pinta_apenas_colunas_a_ate_f(tmp_path):
    headers = HEADERS_COMPLETOS + ["EXTRA1", "EXTRA2", "NAO_DEVE_PINTAR"]
    caminho = _criar_planilha(
        tmp_path, headers, [["123", "C1", "DR. FULANO", "14082026", "x", "y", "z"]]
    )
    registrar_status(caminho, row_index=2, status="sucesso")

    wb = openpyxl.load_workbook(caminho)
    sheet = wb.active
    for col in range(1, 7):
        assert sheet.cell(row=2, column=col).fill.start_color.rgb == COR_SUCESSO
    assert sheet.cell(row=2, column=7).fill.start_color.rgb != COR_SUCESSO
    wb.close()
