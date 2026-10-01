from pathlib import Path
import openpyxl
import pytest
from database.excel_reader import carregar_fila_pacientes, registrar_status, verificar_colunas


@pytest.fixture
def planilha_teste(tmp_path: Path) -> Path:
    arquivo = tmp_path / "pacientes_teste.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fila"

    # Cabeçalho na linha 1
    ws.append(["REG", "CONSULTA", "MÉDICO", "DATA DE VALIDADE", "UNIDADE"])
    # Linha 2: Completa
    ws.append(["1001", "CONSULTA EM OFTALMOLOGIA", "DR. ROBERTO", "15092026", "HOSPITAL OFTALMO PE"])
    # Linha 3: Sem médico e sem unidade na linha (deve herdar da memória se mesclado/arrastado)
    ws.append(["1002", "CONSULTA EM OFTALMOLOGIA", None, "15092026", None])

    wb.save(arquivo)
    wb.close()
    return arquivo


def test_verificar_colunas(planilha_teste: Path):
    faltam = verificar_colunas(planilha_teste)
    assert faltam == []


def test_carregar_fila_com_arraste_memoria(planilha_teste: Path):
    fila = carregar_fila_pacientes(planilha_teste)
    assert len(fila) == 2
    assert fila[0].reg == "1001"
    assert fila[0].unidade == "HOSPITAL OFTALMO PE"
    assert fila[1].reg == "1002"
    # Paciente 2 herdou a unidade da linha anterior por arraste de memória
    assert fila[1].unidade == "HOSPITAL OFTALMO PE"


def test_carregar_fila_sem_unidade_lanca_erro(tmp_path: Path):
    arquivo = tmp_path / "sem_unidade.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["REG", "CONSULTA", "MÉDICO", "DATA DE VALIDADE", "UNIDADE"])
    ws.append(["2001", "CONSULTA", "DR. SILVA", "15092026", None])
    wb.save(arquivo)
    wb.close()

    # Como a primeira linha não tem unidade e não há unidade_global, deve falhar estritamente
    with pytest.raises(ValueError, match="Unidade hospitalar não identificada"):
        carregar_fila_pacientes(arquivo)


def test_registrar_status_pintura(planilha_teste: Path):
    ok = registrar_status(planilha_teste, row_index=2, status="sucesso")
    assert ok is True

    wb = openpyxl.load_workbook(planilha_teste)
    ws = wb.active
    fill = ws.cell(row=2, column=1).fill
    assert fill is not None
    # Deve estar preenchido
    assert fill.fill_type == "solid"
    wb.close()
